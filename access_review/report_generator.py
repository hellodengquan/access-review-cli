from datetime import datetime
from typing import List, Dict, Any, Optional
from pathlib import Path
import csv
import json
from collections import defaultdict

from .models import UserPermission, ReviewRecord, AnomalyReport, ReviewCycle, AnomalyType
from .storage import Storage
from .dtutils import utcnow


try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class ReportGenerator:
    def __init__(self, storage: Storage):
        self.storage = storage

    def generate_review_summary(self, cycle: Optional[ReviewCycle] = None) -> Dict[str, Any]:
        permissions = self.storage.list_permissions()
        reviews = self.storage.list_reviews()
        anomalies = self.storage.list_anomalies(unresolved_only=True)

        if cycle:
            reviews = [
                r for r in reviews
                if cycle.start_date <= r.review_date <= (cycle.end_date or utcnow())
            ]

        total_perms = len(permissions)
        active_perms = [p for p in permissions if p.status == "active"]
        reviewed_perm_ids = set(r.permission_id for r in reviews)
        reviewed_perms = [p for p in active_perms if p.id in reviewed_perm_ids]
        unreviewed_perms = [p for p in active_perms if p.id not in reviewed_perm_ids]

        approved = [r for r in reviews if r.result == "approved"]
        to_revoke = [r for r in reviews if r.result == "revoke_recommended"]
        escalated = [r for r in reviews if r.result == "escalated"]

        by_department: Dict[str, Dict[str, Any]] = defaultdict(lambda: {
            "total": 0, "reviewed": 0, "approved": 0, "to_revoke": 0, "anomalies": 0
        })

        perm_reviews_map = defaultdict(list)
        for r in reviews:
            perm_reviews_map[r.permission_id].append(r)

        for perm in active_perms:
            dept = perm.department
            by_department[dept]["total"] += 1
            if perm.id in reviewed_perm_ids:
                by_department[dept]["reviewed"] += 1
                perm_reviews = perm_reviews_map.get(perm.id, [])
                if any(r.result == "approved" for r in perm_reviews):
                    by_department[dept]["approved"] += 1
                if any(r.result == "revoke_recommended" for r in perm_reviews):
                    by_department[dept]["to_revoke"] += 1

        anomaly_perm_ids = set(a.permission_id for a in anomalies)
        for perm in active_perms:
            if perm.id in anomaly_perm_ids:
                by_department[perm.department]["anomalies"] += 1

        return {
            "generated_at": utcnow().isoformat(),
            "cycle": cycle.model_dump(mode="json") if cycle else None,
            "summary": {
                "total_permissions": total_perms,
                "active_permissions": len(active_perms),
                "reviewed_count": len(reviewed_perms),
                "unreviewed_count": len(unreviewed_perms),
                "review_progress": f"{len(reviewed_perms)}/{len(active_perms)}" if active_perms else "0/0",
                "review_progress_pct": round(len(reviewed_perms) / len(active_perms) * 100, 1) if active_perms else 0,
                "approved_count": len(approved),
                "revoke_recommended_count": len(to_revoke),
                "escalated_count": len(escalated),
                "total_anomalies": len(anomalies),
            },
            "by_department": dict(by_department),
            "anomalies": [a.model_dump(mode="json") for a in anomalies],
            "to_revoke_list": [
                {
                    "permission": self.storage.load_permission(r.permission_id).model_dump(mode="json")
                    if self.storage.load_permission(r.permission_id) else None,
                    "review": r.model_dump(mode="json"),
                }
                for r in to_revoke
            ],
        }

    def generate_revoke_list(self) -> List[Dict[str, Any]]:
        reviews = self.storage.list_reviews()
        to_revoke_reviews = [r for r in reviews if r.result == "revoke_recommended"]

        seen_permissions = set()
        result = []
        for review in to_revoke_reviews:
            perm = self.storage.load_permission(review.permission_id)
            if perm and perm.status != "revoked" and perm.id not in seen_permissions:
                seen_permissions.add(perm.id)
                result.append({
                    "permission_id": perm.id,
                    "username": perm.username,
                    "email": perm.email,
                    "department": perm.department,
                    "role": perm.role,
                    "resource": perm.resource,
                    "permission_level": perm.permission_level,
                    "granted_date": perm.granted_date.isoformat(),
                    "reviewer": review.reviewer,
                    "review_date": review.review_date.isoformat(),
                    "review_comments": review.comments or "",
                })
        return result

    def export_to_csv(self, data: List[Dict[str, Any]], output_path: Path, default_headers: Optional[List[str]] = None) -> None:
        if not data:
            if default_headers:
                with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.DictWriter(f, fieldnames=default_headers)
                    writer.writeheader()
            return

        fieldnames = list(data[0].keys())
        with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(data)

    def export_revoke_list_csv(self, output_path: Path) -> int:
        revoke_list = self.generate_revoke_list()
        default_headers = [
            "permission_id", "username", "email", "department", "role",
            "resource", "permission_level", "granted_date", "reviewer",
            "review_date", "review_comments"
        ]
        self.export_to_csv(revoke_list, output_path, default_headers=default_headers)
        return len(revoke_list)

    def export_permission_list_csv(self, output_path: Path, status: Optional[str] = None) -> int:
        permissions = self.storage.list_permissions()
        if status:
            permissions = [p for p in permissions if p.status == status]

        data = [
            {
                "id": p.id,
                "username": p.username,
                "email": p.email,
                "department": p.department,
                "role": p.role,
                "resource": p.resource,
                "permission_level": p.permission_level,
                "granted_date": p.granted_date.isoformat(),
                "last_used_date": p.last_used_date.isoformat() if p.last_used_date else "",
                "expiry_date": p.expiry_date.isoformat() if p.expiry_date else "",
                "status": p.status.value if hasattr(p.status, 'value') else str(p.status),
                "granted_by": p.granted_by or "",
                "description": p.description or "",
                "tags": ",".join(p.tags),
            }
            for p in permissions
        ]
        default_headers = [
            "id", "username", "email", "department", "role",
            "resource", "permission_level", "granted_date",
            "last_used_date", "expiry_date", "status",
            "granted_by", "description", "tags"
        ]
        self.export_to_csv(data, output_path, default_headers=default_headers)
        return len(data)

    def generate_anomaly_report(self, output_path: Optional[Path] = None) -> Dict[str, Any]:
        anomalies = self.storage.list_anomalies(unresolved_only=True)
        permissions = self.storage.list_permissions()
        perm_map = {p.id: p for p in permissions}

        by_type: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        for anomaly in anomalies:
            perm = perm_map.get(anomaly.permission_id)
            by_type[anomaly.anomaly_type.value].append({
                "anomaly": anomaly.model_dump(mode="json"),
                "permission": perm.model_dump(mode="json") if perm else None,
            })

        report = {
            "generated_at": utcnow().isoformat(),
            "total_unresolved": len(anomalies),
            "by_type": {k: v for k, v in by_type.items()},
        }

        if output_path:
            with open(output_path, "w", encoding="utf-8") as f:
                json.dump(report, f, indent=2, ensure_ascii=False)

        return report

    def generate_user_report(self, username: str) -> Dict[str, Any]:
        permissions = self.storage.list_permissions()
        user_perms = [p for p in permissions if p.username.lower() == username.lower()]
        reviews = []
        for perm in user_perms:
            reviews.extend(self.storage.get_reviews_for_permission(perm.id))
        anomalies = []
        for perm in user_perms:
            anomalies.extend(self.storage.get_anomalies_for_permission(perm.id))

        return {
            "generated_at": utcnow().isoformat(),
            "username": username,
            "permissions": [p.model_dump(mode="json") for p in user_perms],
            "reviews": [r.model_dump(mode="json") for r in reviews],
            "anomalies": [a.model_dump(mode="json") for a in anomalies],
            "permission_count": len(user_perms),
            "active_count": len([p for p in user_perms if p.status == "active"]),
            "anomaly_count": len([a for a in anomalies if not a.resolved]),
        }

    def generate_aggregate_report(self, group_by: str = "department") -> Dict[str, Any]:
        valid_group_by = {"user", "role", "department"}
        if group_by not in valid_group_by:
            raise ValueError(
                f"group_by 必须是 {valid_group_by} 之一，当前值: {group_by}"
            )

        permissions = self.storage.list_permissions()
        anomalies = self.storage.list_anomalies(unresolved_only=True)

        anomaly_perm_ids = set(a.permission_id for a in anomalies)

        groups: Dict[str, Dict[str, Any]] = defaultdict(lambda: {
            "group_key": "",
            "permission_count": 0,
            "active_count": 0,
            "expired_count": 0,
            "revoked_count": 0,
            "pending_count": 0,
            "anomaly_count": 0,
            "high_risk_count": 0,
            "resources": set(),
            "roles": set(),
            "users": set(),
            "departments": set(),
            "permissions": [],
        })

        for perm in permissions:
            if group_by == "user":
                key = perm.username
            elif group_by == "role":
                key = perm.role
            else:
                key = perm.department

            g = groups[key]
            g["group_key"] = key
            g["permission_count"] += 1
            g["users"].add(perm.username)
            g["departments"].add(perm.department)
            g["roles"].add(perm.role)
            g["resources"].add(perm.resource)

            status_val = perm.status.value if hasattr(perm.status, 'value') else str(perm.status)
            if status_val == "active":
                g["active_count"] += 1
            elif status_val == "expired":
                g["expired_count"] += 1
            elif status_val == "revoked":
                g["revoked_count"] += 1
            elif status_val == "pending_review":
                g["pending_count"] += 1

            if perm.id in anomaly_perm_ids:
                g["anomaly_count"] += 1
                perm_anomalies = [a for a in anomalies if a.permission_id == perm.id]
                for a in perm_anomalies:
                    sev = a.severity.value if hasattr(a.severity, 'value') else str(a.severity)
                    if sev in ("high", "critical"):
                        g["high_risk_count"] += 1
                        break

            g["permissions"].append({
                "id": perm.id,
                "username": perm.username,
                "department": perm.department,
                "role": perm.role,
                "resource": perm.resource,
                "permission_level": perm.permission_level,
                "status": status_val,
                "granted_date": perm.granted_date.isoformat(),
                "last_used_date": perm.last_used_date.isoformat() if perm.last_used_date else None,
            })

        group_list = []
        for key, g in sorted(groups.items(), key=lambda x: -x[1]["permission_count"]):
            group_list.append({
                "group_key": g["group_key"],
                "permission_count": g["permission_count"],
                "active_count": g["active_count"],
                "expired_count": g["expired_count"],
                "revoked_count": g["revoked_count"],
                "pending_count": g["pending_count"],
                "anomaly_count": g["anomaly_count"],
                "high_risk_count": g["high_risk_count"],
                "user_count": len(g["users"]),
                "department_count": len(g["departments"]),
                "role_count": len(g["roles"]),
                "resource_count": len(g["resources"]),
                "permissions": g["permissions"],
            })

        return {
            "generated_at": utcnow().isoformat(),
            "group_by": group_by,
            "total_groups": len(group_list),
            "total_permissions": len(permissions),
            "total_anomalies": len(anomalies),
            "groups": group_list,
        }

    def export_aggregate_csv(self, report: Dict[str, Any], output_path: Path) -> int:
        group_by = report.get("group_by", "department")
        groups = report.get("groups", [])

        label_map = {
            "user": "用户",
            "role": "角色",
            "department": "部门",
        }
        group_label = label_map.get(group_by, group_by)

        summary_rows = []
        for g in groups:
            summary_rows.append({
                group_label: g["group_key"],
                "权限总数": g["permission_count"],
                "活跃权限": g["active_count"],
                "已过期": g["expired_count"],
                "已收回": g["revoked_count"],
                "待复核": g["pending_count"],
                "异常数": g["anomaly_count"],
                "高风险数": g["high_risk_count"],
                "涉及用户数": g["user_count"],
                "涉及部门数": g["department_count"],
                "涉及角色数": g["role_count"],
                "涉及资源数": g["resource_count"],
            })

        summary_headers = list(summary_rows[0].keys()) if summary_rows else [group_label, "权限总数"]
        detail_rows = []
        for g in groups:
            for idx, p in enumerate(g.get("permissions", []), start=1):
                row = {
                    group_label: g["group_key"],
                    "序号": idx,
                    "权限ID": p["id"],
                    "用户名": p["username"],
                    "部门": p["department"],
                    "角色": p["role"],
                    "资源": p["resource"],
                    "权限级别": p["permission_level"],
                    "状态": p["status"],
                    "授予日期": p["granted_date"].split("T")[0] if p["granted_date"] else "",
                    "最后使用": p["last_used_date"].split("T")[0] if p["last_used_date"] else "",
                }
                detail_rows.append(row)

        detail_headers = list(detail_rows[0].keys()) if detail_rows else []

        if output_path.suffix.lower() != ".csv":
            output_path = output_path.with_suffix(".csv")

        with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
            if summary_rows:
                writer = csv.DictWriter(f, fieldnames=summary_headers)
                f.write(f"# 聚合视图 - 按{group_label}分组 (概览)\n")
                writer.writeheader()
                writer.writerows(summary_rows)
                f.write("\n")

            if detail_rows:
                writer = csv.DictWriter(f, fieldnames=detail_headers)
                f.write(f"# 聚合视图 - 按{group_label}分组 (明细)\n")
                writer.writeheader()
                writer.writerows(detail_rows)

        return len(summary_rows) + len(detail_rows)

    def export_aggregate_excel(self, report: Dict[str, Any], output_path: Path) -> int:
        if not HAS_OPENPYXL:
            raise ImportError(
                "缺少 openpyxl 依赖，无法导出 Excel。"
                "请执行: pip install openpyxl"
            )

        group_by = report.get("group_by", "department")
        groups = report.get("groups", [])

        label_map = {
            "user": "用户",
            "role": "角色",
            "department": "部门",
        }
        group_label = label_map.get(group_by, group_by)

        wb = Workbook()

        ws_summary = wb.active
        ws_summary.title = f"概览-{group_label}"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        summary_headers = [
            group_label, "权限总数", "活跃权限", "已过期", "已收回", "待复核",
            "异常数", "高风险数", "涉及用户数", "涉及部门数", "涉及角色数", "涉及资源数",
        ]

        for col_idx, header in enumerate(summary_headers, start=1):
            cell = ws_summary.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        for row_idx, g in enumerate(groups, start=2):
            row_data = [
                g["group_key"],
                g["permission_count"],
                g["active_count"],
                g["expired_count"],
                g["revoked_count"],
                g["pending_count"],
                g["anomaly_count"],
                g["high_risk_count"],
                g["user_count"],
                g["department_count"],
                g["role_count"],
                g["resource_count"],
            ]
            for col_idx, val in enumerate(row_data, start=1):
                ws_summary.cell(row=row_idx, column=col_idx, value=val)

        for col_idx in range(1, len(summary_headers) + 1):
            ws_summary.column_dimensions[chr(64 + col_idx)].width = 14
        ws_summary.column_dimensions["A"].width = 20

        ws_detail = wb.create_sheet(title=f"明细-{group_label}")
        detail_headers = [
            group_label, "序号", "权限ID", "用户名", "部门", "角色", "资源",
            "权限级别", "状态", "授予日期", "最后使用",
        ]
        for col_idx, header in enumerate(detail_headers, start=1):
            cell = ws_detail.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        row_idx = 2
        for g in groups:
            for idx, p in enumerate(g.get("permissions", []), start=1):
                row_data = [
                    g["group_key"],
                    idx,
                    p["id"],
                    p["username"],
                    p["department"],
                    p["role"],
                    p["resource"],
                    p["permission_level"],
                    p["status"],
                    p["granted_date"].split("T")[0] if p["granted_date"] else "",
                    p["last_used_date"].split("T")[0] if p["last_used_date"] else "",
                ]
                for col_idx, val in enumerate(row_data, start=1):
                    ws_detail.cell(row=row_idx, column=col_idx, value=val)
                row_idx += 1

        ws_detail_widths = [20, 6, 16, 12, 12, 12, 22, 10, 10, 12, 12]
        for col_idx, width in enumerate(ws_detail_widths, start=1):
            ws_detail.column_dimensions[chr(64 + col_idx)].width = width

        ws_meta = wb.create_sheet(title="元信息")
        meta_rows = [
            ("生成时间", report.get("generated_at", "")),
            ("聚合维度", group_by),
            ("分组总数", report.get("total_groups", 0)),
            ("总权限数", report.get("total_permissions", 0)),
            ("未解决异常数", report.get("total_anomalies", 0)),
        ]
        for row_idx, (key, val) in enumerate(meta_rows, start=1):
            ws_meta.cell(row=row_idx, column=1, value=key).font = header_font
            ws_meta.cell(row=row_idx, column=1).fill = header_fill
            ws_meta.cell(row=row_idx, column=2, value=val)
        ws_meta.column_dimensions["A"].width = 18
        ws_meta.column_dimensions["B"].width = 40

        if output_path.suffix.lower() not in (".xlsx", ".xlsm"):
            output_path = output_path.with_suffix(".xlsx")
        wb.save(output_path)

        return len(groups) + sum(len(g.get("permissions", [])) for g in groups)
