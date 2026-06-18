#!/usr/bin/env python3
import os
import shutil
import json
import glob
import sys

os.environ['ACCESS_REVIEW_OPERATOR'] = 'auditor_zhang'
os.environ['ACCESS_REVIEW_HOME'] = '/tmp/test_final_6'
shutil.rmtree('/tmp/test_final_6', ignore_errors=True)

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from typer.testing import CliRunner
from access_review.cli import app

runner = CliRunner()
env = {'ACCESS_REVIEW_OPERATOR': 'auditor_zhang', 'ACCESS_REVIEW_HOME': '/tmp/test_final_6'}

def run_cmd(args, expected=0, print_output=True):
    result = runner.invoke(app, args, env=env)
    if print_output:
        print(result.output)
    print(f'EXIT: {result.exit_code} (expect {expected})')
    if result.exception and not isinstance(result.exception, SystemExit):
        import traceback
        traceback.print_exception(type(result.exception), result.exception, result.exception.__traceback__)
    assert result.exit_code == expected, f'失败: {args}, exit={result.exit_code}'
    return result


def main():
    print('=' * 70)
    print('测试1: 导入跨OU权限 (email归一化+username后缀清理)')
    print('=' * 70)
    run_cmd(['permission', 'import', '/tmp/test_cross_ou.json'])

    print()
    print('=' * 70)
    print('测试2: 统计验证 (4条输入 → 合并后预期剩余3条)')
    print('=' * 70)
    run_cmd(['stats'])

    print()
    print('=' * 70)
    print('测试3: 聚合视图 - 按部门')
    print('=' * 70)
    run_cmd(['report', 'aggregate', '-g', 'department'])

    print()
    print('=' * 70)
    print('测试4: 聚合视图 - 按角色 + 明细')
    print('=' * 70)
    run_cmd(['report', 'aggregate', '-g', 'role', '-d', '-n', '10'])

    print()
    print('=' * 70)
    print('测试5: CSV导出')
    print('=' * 70)
    run_cmd(['report', 'aggregate', '-g', 'user', '-f', 'csv', '-o', '/tmp/r_user.csv'])
    with open('/tmp/r_user.csv', encoding='utf-8-sig') as f:
        content = f.read()
    print('--- CSV 前10行 ---')
    for i, line in enumerate(content.splitlines()[:10], 1):
        print(f'{i:2}. {line}')
    print('-----------------')

    print()
    print('=' * 70)
    print('测试6: Excel导出')
    print('=' * 70)
    run_cmd(['report', 'aggregate', '-g', 'department', '-f', 'excel', '-o', '/tmp/r_dept.xlsx'])
    size = os.path.getsize('/tmp/r_dept.xlsx')
    print(f'文件大小: {size} 字节')
    try:
        from openpyxl import load_workbook
        wb = load_workbook('/tmp/r_dept.xlsx')
        print(f'Sheet列表: {wb.sheetnames}')
        ws = wb[wb.sheetnames[0]]
        print(f'概览sheet: {ws.max_row} 行 x {ws.max_column} 列')
        for row in ws.iter_rows(min_row=1, max_row=min(4, ws.max_row), values_only=True):
            print('  | '.join(str(c) for c in row if c is not None))
    except Exception as e:
        print(f'读Excel出错: {e}')

    print()
    print('=' * 70)
    print('测试7: 审计日志列表')
    print('=' * 70)
    r = run_cmd(['audit', 'list', '-n', '100'])

    print()
    print('=' * 70)
    print('测试8: 检查审计日志原始内容')
    print('=' * 70)
    files = sorted(glob.glob('/tmp/test_final_6/audit/*.jsonl'))
    for fp in files:
        print(f'日志文件: {fp}')
        with open(fp) as fh:
            lines = [l.strip() for l in fh.readlines() if l.strip()]
        print(f'  总行数: {len(lines)}')
        sample_indices = [0, 1] if len(lines) >= 2 else list(range(len(lines)))
        for idx in sample_indices:
            e = json.loads(lines[idx])
            print(f'  --- Entry #{idx+1} (action={e.get("action")}, status={e.get("status")})')
            keys_show = ['timestamp', 'action', 'status', 'operator', 'hostname', 'command', 'params', 'error', 'duration_ms']
            for k in keys_show:
                if k in e and e[k] is not None:
                    print(f'    {k}: {e[k]}')

    print()
    print('=' * 70)
    print('✅ 所有测试通过！')
    print('=' * 70)


if __name__ == '__main__':
    main()
